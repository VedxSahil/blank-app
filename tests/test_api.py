"""API tests with a TEST FIXTURE corpus."""

from fastapi.testclient import TestClient

from pereport.api import create_app
from pereport.chunker import Chunk

CHUNKS = [
    Chunk("fx:p1:c0", "fx", 1, "FEES", "2024-01-01", "F", "text", "The management fee is 2% of committed capital."),
    Chunk("fx:p2:c0", "fx", 2, "NAV", "2024-01-01", "F", "text", "The net asset value was 500 million at 30 June 2023."),
]  # fmt: skip


def _client(threshold: float | None = 0.2, chunks: list[Chunk] | None = None) -> TestClient:
    return TestClient(create_app(CHUNKS if chunks is None else chunks, threshold))


def test_ask_answer_and_refusal() -> None:
    c = _client()
    ok = c.post("/api/ask", json={"question": "What is the management fee?"}).json()
    assert not ok["refused"] and ok["verified"] is True
    assert ok["text"] == "The management fee is 2% of committed capital [fx, p.1]."
    assert [s["chunk_id"] for s in ok["sources"]] == ["fx:p1:c0"]
    no = c.post("/api/ask", json={"question": "What is the clawback?"}).json()
    assert no["refused"] and no["reason"] == "no_evidence" and no["sources"] == []


def test_uncalibrated_or_empty_returns_503_not_a_guess() -> None:
    assert _client(threshold=None).post("/api/ask", json={"question": "fee"}).status_code == 503
    assert _client(chunks=[]).post("/api/ask", json={"question": "fee"}).status_code == 503
    assert _client(threshold=None).get("/api/health").json() == {
        "ok": True, "chunks": 2, "calibrated": False, "banner": None,
    }  # fmt: skip


def test_empty_question_rejected() -> None:
    assert _client().post("/api/ask", json={"question": "  "}).status_code == 422


def test_report_chunk_documents_eval() -> None:
    c = _client()
    rep = c.post("/api/report").json()
    assert rep["total"] == 17 and 0 <= rep["disclosed"] <= rep["total"]
    flat = [ln for s in rep["sections"] for ln in s["lines"]]
    assert all(ln["chunk_ids"] for ln in flat if ln["disclosed"])
    assert all(
        ln["text"] == "Not disclosed in provided documents" for ln in flat if not ln["disclosed"]
    )
    assert c.get("/api/chunks/fx:p2:c0").json()["page"] == 2
    assert c.get("/api/chunks/nope").status_code == 404
    docs = c.get("/api/documents").json()["documents"]
    assert docs == [{"doc": "fx", "chunks": 2, "pages": 2, "date": "2024-01-01", "uploaded": False}]
    assert c.get("/api/eval").json() == {"available": False}


def test_banner_is_reported_and_app_from_disk_reads_env(tmp_path, monkeypatch) -> None:  # type: ignore[no-untyped-def]
    from pereport.api import app_from_disk
    from pereport.store import save_chunks

    assert (
        TestClient(create_app(CHUNKS, 0.2, banner="SAMPLE")).get("/api/health").json()["banner"]
        == "SAMPLE"
    )
    save_chunks(tmp_path / "chunks" / "d.jsonl", CHUNKS)
    (tmp_path / "ev").mkdir()
    (tmp_path / "ev" / "gate.json").write_text('{"bm25": 0.2}')
    monkeypatch.setenv("PEREPORT_CHUNKS", str(tmp_path / "chunks"))
    monkeypatch.setenv("PEREPORT_EVAL", str(tmp_path / "ev"))
    monkeypatch.setenv("PEREPORT_BANNER", "FROM ENV")
    health = TestClient(app_from_disk()).get("/api/health").json()
    assert health == {"ok": True, "chunks": 2, "calibrated": True, "banner": "FROM ENV"}


def _pdf_b64(text_lines: list[str]) -> str:
    import base64

    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    for i, line in enumerate(text_lines):
        page.insert_text((72, 80 + 20 * i), line, fontsize=11)
    return base64.b64encode(doc.tobytes()).decode()


def test_upload_pdf_then_ask_it_scoped_and_unscoped() -> None:
    c = _client()
    b64 = _pdf_b64(["The performance fee is 15 percent of profits above the benchmark."])
    r = c.post("/api/upload", json={"filename": "My Fund (v2).pdf", "data_base64": b64})
    assert r.status_code == 200, r.text
    up = r.json()
    assert (
        up["doc"] == "My_Fund_v2"
        and up["pages"] == 1
        and up["chunks"] >= 1
        and up["kind"] == "pages"
    )
    assert any("refusal threshold" in n for n in up["notes"])
    docs = {d["doc"]: d for d in c.get("/api/documents").json()["documents"]}
    assert docs["My_Fund_v2"]["uploaded"] is True and docs["fx"]["uploaded"] is False
    # scoped to the new document: answered from it, with a citation to its page
    a = c.post(
        "/api/ask", json={"question": "What is the performance fee?", "doc": "My_Fund_v2"}
    ).json()
    assert not a["refused"] and "[My_Fund_v2, p.1]" in a["text"] and a["verified"] is True
    # scoped to the original document: the uploaded text is not visible there
    b = c.post("/api/ask", json={"question": "What is the performance fee?", "doc": "fx"}).json()
    assert b["refused"]
    assert c.post("/api/ask", json={"question": "x y", "doc": "nope"}).status_code == 404
    # a second upload with the same name gets a distinct document id
    r2 = c.post("/api/upload", json={"filename": "My Fund (v2).pdf", "data_base64": b64}).json()
    assert r2["doc"] == "My_Fund_v2_2"


def test_upload_xlsx_indexes_text_and_says_numbers_are_dropped(tmp_path) -> None:  # type: ignore[no-untyped-def]
    import base64

    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Fees"
    ws["B2"], ws["C2"] = "Field", "Definition"
    ws["B3"], ws["C3"] = "Carry", "The GP share of profits."
    ws["D3"] = 987654
    path = tmp_path / "t.xlsx"
    wb.save(path)
    c = _client()
    r = c.post(
        "/api/upload",
        json={"filename": "t.xlsx", "data_base64": base64.b64encode(path.read_bytes()).decode()},
    ).json()
    assert r["kind"] == "sheets" and any("numbers and dates" in n for n in r["notes"])
    chunk = c.get(f"/api/chunks/{r['doc']}:p1:c0").json()
    assert "987654" not in chunk["text"]


def test_upload_rejections() -> None:
    import base64

    c = _client()
    ok_b64 = _pdf_b64(["Some text here."])
    assert (
        c.post("/api/upload", json={"filename": "a.docx", "data_base64": ok_b64}).status_code == 415
    )
    assert (
        c.post("/api/upload", json={"filename": "a.pdf", "data_base64": "***"}).status_code == 422
    )
    junk = base64.b64encode(b"this is not a pdf").decode()
    assert c.post("/api/upload", json={"filename": "a.pdf", "data_base64": junk}).status_code == 422
    blank = _pdf_b64([])  # a page with no text: a scan would look like this
    r = c.post("/api/upload", json={"filename": "blank.pdf", "data_base64": blank})
    assert r.status_code == 422 and "OCR" in r.json()["detail"]
    huge = "A" * (26 * 1024 * 1024 * 4 // 3)
    assert c.post("/api/upload", json={"filename": "a.pdf", "data_base64": huge}).status_code == 413
    # nothing from the rejected files was added
    assert [d["doc"] for d in c.get("/api/documents").json()["documents"]] == ["fx"]
