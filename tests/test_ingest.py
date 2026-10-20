"""Ingest tests on a PDF generated here from TEST FIXTURE text (not a real document)."""

from pathlib import Path

import pymupdf

from pereport.chunker import chunk_pages
from pereport.ingest import parse_pdf
from pereport.store import load_chunks, save_chunks


def _make_pdf(path: Path) -> None:
    doc = pymupdf.open()
    p1 = doc.new_page()
    p1.insert_text((72, 80), "1. FEES", fontsize=14)
    p1.insert_text((72, 110), "The management fee is 2 percent of committed capital.", fontsize=10)
    p2 = doc.new_page()
    p2.insert_text((72, 80), "2. TABLE", fontsize=14)
    # a ruled 2x3 table so find_tables can see it
    x = [72, 172, 272, 372]
    y = [120, 150, 180]
    for yy in y:
        p2.draw_line((x[0], yy), (x[-1], yy))
    for xx in x:
        p2.draw_line((xx, y[0]), (xx, y[-1]))
    cells = [["Year", "Called", "Paid"], ["2021", "100", "40"]]
    for r, row in enumerate(cells):
        for c, val in enumerate(row):
            p2.insert_text((x[c] + 4, y[r] + 20), val, fontsize=10)
    doc.new_page()  # blank page, no images -> not flagged for OCR
    doc.save(path)


def test_parse_keeps_page_numbers_and_table_rows(tmp_path: Path) -> None:
    pdf = tmp_path / "fx.pdf"
    _make_pdf(pdf)
    parsed = parse_pdf(pdf)
    assert [p.page for p in parsed.pages] == [1, 2, 3]
    assert "management fee is 2 percent" in parsed.pages[0].text
    assert "2021  100  40" in parsed.pages[1].text  # one table row per line
    assert parsed.ocr_pages == []


def test_pdf_to_chunks_roundtrip_through_store(tmp_path: Path) -> None:
    pdf = tmp_path / "fx.pdf"
    _make_pdf(pdf)
    chunks = chunk_pages(parse_pdf(pdf).pages, doc="fx", date="2024-01-01", fund="F")
    fee = next(c for c in chunks if "management fee" in c.text)
    assert fee.page == 1 and fee.section == "1. FEES"
    table = next(c for c in chunks if c.kind == "table")
    assert table.page == 2 and "2021  100  40" in table.text
    save_chunks(tmp_path / "c" / "fx.jsonl", chunks)
    assert load_chunks(tmp_path / "c" / "fx.jsonl") == chunks


def test_merge_continuations_empty_first_cell_and_sentence_style() -> None:
    from pereport.ingest import merge_continuations

    # empty first cell always continues the row above
    assert merge_continuations([["Role", "Entity"], ["Fund", "Alpha"], ["", "Beta"]]) == [
        ["Role", "Entity"],
        ["Fund", "Alpha Beta"],
    ]
    # sentence-style table: a row whose text has not reached a full stop continues on the next line
    grid = [["Term", "Summary"], ["Alternate", "May be organized in any"], ["Vehicles", "jurisdiction."],
            ["Fee", "Two percent."], ["Carry", "Twenty percent."]]  # fmt: skip
    assert merge_continuations(grid)[1:] == [
        ["Alternate Vehicles", "May be organized in any jurisdiction."],
        ["Fee", "Two percent."],
        ["Carry", "Twenty percent."],
    ]
    # address-style table (no full stops) is NOT merged by the sentence rule
    addr = [["Role", "Address"], ["Fund", "1 Quay"], ["Sponsor", "2 Tower"]]
    assert merge_continuations(addr) == addr


def test_group_by_rules_merges_lines_between_the_same_rules() -> None:
    from pereport.ingest import group_by_rules

    # Rules at y=8, 20, 40 give buckets: y<8 -> 0, 8..20 -> 1, 20..40 -> 2.
    # Header (y=5) is bucket 0; "A" (y=10) and its wrapped line (y=15) are bucket 1; "B" (y=30) is bucket 2.
    grid = [["Head", "Cell"], ["A", "one two"], ["", "three"], ["B", "four"]]
    out = group_by_rules(grid, [5.0, 10.0, 15.0, 30.0], [8.0, 20.0, 40.0])
    assert out == [["Head", "Cell"], ["A", "one two three"], ["B", "four"]]


def test_sample_ppm_tables_keep_words_and_rows() -> None:
    """Regression on the user-supplied replica PPM (skipped if the file is absent)."""
    from pathlib import Path

    import pytest

    pdf = Path("data/sample/Aurelian_Peak_Fund_Sample_PPM.pdf")
    if not pdf.exists():
        pytest.skip("sample PDF not present")
    pages = parse_pdf(pdf).pages
    assert len(pages) == 20
    summary = pages[7].text  # PDF page 8
    assert "Open-ended Zandrian public company" in summary  # words were being glued together
    assert "Management Fee  2% per annum of Class A Net Asset Value" in summary
    assert (
        "Investment Manager  Aurelian Asset Management Ltd; discretionary" in summary
    )  # wrapped label
    assert "Legal Advisors (Zandria)  Isolde Marchetti, Attorney" in pages[1].text


def _make_workbook(path: Path) -> None:
    import datetime

    import openpyxl
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Template"
    ws["B2"] = "A. Statement for LP #9"
    ws["B3"] = "A.1 Reconciliation"
    ws["E3"] = "LP share"
    ws["F3"] = "Total fund"
    ws["B4"] = "Fees"  # bold, no values: a sub-heading
    ws["B4"].font = Font(bold=True)
    ws["B5"] = "Management Fees"
    ws["E5"], ws["F5"] = 123456, 987654
    ws["B6"] = "Total Fees"
    ws["E6"] = "(=)"
    ws["F6"] = 111111
    ws["B7"] = "Footnote: an amount in this row needs a description, as explained here."
    ws["B8"] = "(=)"  # bare formula marker: dropped
    ws["H2"] = datetime.date(2022, 2, 23)  # date cell: never indexed
    ws["B10"] = "Reconciliation for Carry"  # bold heading whose first item shares the row
    ws["B10"].font = Font(bold=True)
    ws["C10"] = "Starting period balance"
    ws["E10"] = 5
    ws["C11"] = "Carry accrued"
    wb.create_sheet("Defs")
    d = wb["Defs"]
    d["B2"], d["C2"], d["D2"] = "Section", "Field", "Definition"
    for i in range(10):
        r = 3 + i
        if i in (0, 5):
            d.cell(r, 2).value = f"S{i} Group"
        d.cell(r, 3).value = f"Field {i}"
        d.cell(r, 4).value = f"Meaning of field {i}."
    wb.save(path)


def test_parse_xlsx_headings_line_items_numbers_and_defs(tmp_path: Path) -> None:
    from pereport.ingest import parse_xlsx

    x = tmp_path / "t.xlsx"
    _make_workbook(x)
    pages = parse_xlsx(x).pages
    assert [p.page for p in pages] == [1, 2]  # one page per sheet
    t = pages[0].text.splitlines()
    assert t[0] == "# Template"
    assert "# A. Statement for LP #9" in t
    # lettered row with extra text cells: heading plus header and column groups
    assert t[t.index("# A.1 Reconciliation") + 1 :][:2] == [
        "Line item  Section",
        "Column groups  LP share / Total fund",
    ]
    assert "# A.1 Reconciliation > Fees" in t  # bold row without values: sub-heading
    assert "Management Fees  (template line item in A.1 Reconciliation > Fees)" in t
    assert "Total Fees  (template line item in A.1 Reconciliation > Fees, calculated row)" in t
    assert any(ln.startswith("Footnote: an amount") for ln in t)  # long sentence stays a paragraph
    # a bold heading with its first item on the same row opens a new block: later items belong to it
    assert "# A.1 Reconciliation > Reconciliation for Carry" in t
    assert (
        "Starting period balance  (template line item in A.1 Reconciliation > Reconciliation for Carry)"
        in t
    )
    assert (
        "Carry accrued  (template line item in A.1 Reconciliation > Reconciliation for Carry)" in t
    )
    joined = pages[0].text
    for bad in ("123456", "987654", "111111", "2022", "(=)"):
        assert bad not in joined  # values, dates and bare markers are not indexed
    # definitions table: sparse first column becomes headings, header repeated after each
    d = pages[1].text.splitlines()
    assert d[d.index("# S0 Group") + 1 :][:2] == [
        "Field  Definition",
        "Field 0  Meaning of field 0.",
    ]
    assert "# S5 Group" in d and d[d.index("# S5 Group") + 1] == "Field  Definition"
    assert "Field 7  Meaning of field 7." in d


def test_xlsx_to_chunks_keeps_sheet_page_section_and_rows(tmp_path: Path) -> None:
    from pereport.ingest import parse_xlsx

    x = tmp_path / "t.xlsx"
    _make_workbook(x)
    chunks = chunk_pages(parse_xlsx(x).pages, doc="t")
    f7 = next(c for c in chunks if "Field 7" in c.text)
    assert (f7.page, f7.section, f7.kind) == (2, "S5 Group", "table")
    assert f7.text == "Field  Definition\nField 7  Meaning of field 7."
    mf = next(c for c in chunks if c.text.startswith("Line item  Section\nManagement Fees"))
    assert mf.page == 1 and mf.section == "A.1 Reconciliation > Fees"


def test_ilpa_template_real_file_regression() -> None:
    import pytest

    from pereport.ingest import parse_xlsx

    f = Path("data/real/ILPA_Reporting_Template_v2.0.xlsx")
    if not f.exists():
        pytest.skip("ILPA template not present")
    pages = parse_xlsx(f).pages
    assert [p.page for p in pages] == [
        1,
        2,
        3,
        4,
        5,
    ]  # Guidance, Template, FoF, Definitions, Revisions
    defs = pages[3].text
    assert "Placement Fees  Fees/costs paid to the investment adviser" in defs
    assert "# A.1 NAV Reconciliation" in defs
    assert "Contributions \u2013 Cash & Non-Cash" in pages[1].text
    import re

    assert not re.search(r"\d{5,}", "\n".join(p.text for p in pages))  # no example figures indexed


def test_label_only_rows_never_read_like_a_value(tmp_path: Path) -> None:
    import openpyxl

    from pereport.ingest import parse_xlsx

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Fund of Funds"
    ws["C6"] = "Parent Fund Size:"  # an empty field in a blank template
    path = tmp_path / "blank.xlsx"
    wb.save(path)
    text = parse_xlsx(path).pages[0].text
    assert "Parent Fund Size:  (template line item in Fund of Funds)" in text
    assert "Parent Fund Size:  Fund of Funds\n" not in text + "\n"
