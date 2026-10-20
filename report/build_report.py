import re, sys, json, subprocess
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT="Times New Roman"
TOC_PAGES = json.load(open("toc.json")) if len(sys.argv)>1 else {}

d=Document()
sec=d.sections[0]
sec.page_height,sec.page_width=Cm(29.7),Cm(21.0)
sec.left_margin=Cm(2.8); sec.right_margin=Cm(2.3); sec.top_margin=sec.bottom_margin=Cm(2.3)
sec.different_first_page_header_footer=True

def setfont(run,size,bold=False,italic=False):
    run.font.name=FONT; run.font.size=Pt(size); run.bold=bold; run.italic=italic
    run._element.rPr.rFonts.set(qn('w:eastAsia'),FONT)
st=d.styles['Normal']; st.font.name=FONT; st.font.size=Pt(12)
st.element.rPr.rFonts.set(qn('w:eastAsia'),FONT)

def para(text="",size=12,bold=False,italic=False,align=WD_ALIGN_PARAGRAPH.JUSTIFY,spacing=1.5,after=6,indent=None,keep=False):
    p=d.add_paragraph(); pf=p.paragraph_format
    pf.alignment=align; pf.line_spacing=spacing; pf.space_after=Pt(after); pf.space_before=Pt(0)
    if indent is not None: pf.left_indent=Cm(indent)
    if keep: pf.keep_with_next=True
    if text: setfont(p.add_run(text),size,bold,italic)
    return p
def pagebreak(): pass

# footer page number
def add_field(run,instr):
    for t,txt in (("begin",None),(None,instr),("end",None)):
        if t:
            e=OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'),t)
        else:
            e=OxmlElement('w:instrText'); e.set(qn('xml:space'),'preserve'); e.text=txt
        run._r.append(e)
fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=fp.add_run(); setfont(r,12); add_field(r,"PAGE")

# ---- title page
for _ in range(2): para("",after=0)
para("A PROJECT REPORT",14,True,align=WD_ALIGN_PARAGRAPH.CENTER,after=12)
para("on",12,align=WD_ALIGN_PARAGRAPH.CENTER,after=12)
para("PEREPORT: A BROWSER-BASED RETRIEVAL-AUGMENTED GENERATION (RAG) FUND DOCUMENT ASSISTANT FOR INVESTOR QUERIES, FEE CALCULATION AND REPORT DRAFTING",14,True,align=WD_ALIGN_PARAGRAPH.CENTER,after=30)
para("Submitted in partial fulfilment of the requirements for the award of the degree of",12,align=WD_ALIGN_PARAGRAPH.CENTER,after=12)
para("[Degree Name]",14,True,align=WD_ALIGN_PARAGRAPH.CENTER,after=30)
para("Submitted by",12,align=WD_ALIGN_PARAGRAPH.CENTER,after=6)
para("[Student Name]",14,True,align=WD_ALIGN_PARAGRAPH.CENTER,after=2)
para("[Roll Number / Enrolment Number]",12,align=WD_ALIGN_PARAGRAPH.CENTER,after=24)
para("Under the guidance of",12,align=WD_ALIGN_PARAGRAPH.CENTER,after=6)
para("[Guide Name]",14,True,align=WD_ALIGN_PARAGRAPH.CENTER,after=30)
para("[Department Name]",12,True,align=WD_ALIGN_PARAGRAPH.CENTER,after=2)
para("[Institution Name]",12,True,align=WD_ALIGN_PARAGRAPH.CENTER,after=2)
para("[Academic Year]",12,align=WD_ALIGN_PARAGRAPH.CENTER,after=0)
pagebreak()

lines=open("content.txt",encoding="utf8").read().split("\n")
heads=[]  # (level,text)
for l in lines:
    if l.startswith("##"): heads.append((2,l[2:].strip()))
    elif l.startswith("#"): heads.append((1,l[1:].strip()))

def toc():
    para("TABLE OF CONTENTS",14,True,align=WD_ALIGN_PARAGRAPH.CENTER,after=10).paragraph_format.page_break_before=True
    for lvl,t in heads:
        if lvl==2 and not re.match(r"\d+\.\d+ ",t): continue
        p=para("",12,spacing=1.15,after=2,align=WD_ALIGN_PARAGRAPH.LEFT,indent=0 if lvl==1 else 0.8)
        pf=p.paragraph_format
        pf.tab_stops.add_tab_stop(Cm(15.29),alignment=2,leader=1)  # right, dots
        setfont(p.add_run(t),12,bold=(lvl==1))
        setfont(p.add_run("\t"+str(TOC_PAGES.get(t,""))),12,bold=(lvl==1))

first_chapter_done=False; toc_done=False
i=0
def add_table(title,rows):
    para(title,12,True,align=WD_ALIGN_PARAGRAPH.CENTER,spacing=1.0,after=4,keep=True)
    t=d.add_table(rows=len(rows),cols=len(rows[0])); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for ri,row in enumerate(rows):
        for ci,val in enumerate(row):
            c=t.cell(ri,ci); c.text=""
            p=c.paragraphs[0]; p.paragraph_format.line_spacing=1.0; p.paragraph_format.space_after=Pt(2)
            p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
            setfont(p.add_run(val),12,bold=(ri==0))
            if ri==0:
                sh=OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:fill'),'D9D9D9'); c._tc.get_or_add_tcPr().append(sh)
    if "A.1" in title:
        t.autofit=False
        for row in t.rows:
            for ci,w in enumerate((1.2,9.0,2.0,3.6)): row.cells[ci].width=Cm(w)
    para("",spacing=1.0,after=6)

while i<len(lines):
    l=lines[i]; i+=1
    if not l.strip(): continue
    if l.startswith("##"):
        para(l[2:].strip(),12,True,align=WD_ALIGN_PARAGRAPH.LEFT,after=4,keep=True)
    elif l.startswith("#"):
        t=l[1:].strip()
        if not toc_done and t=="EXECUTIVE SUMMARY":
            pass
        if t!="EXECUTIVE SUMMARY": pagebreak()
        hp=para(t,14,True,align=WD_ALIGN_PARAGRAPH.CENTER,after=12,keep=True)
        hp.paragraph_format.page_break_before=not t.startswith('APPENDIX B')
        if t.startswith('APPENDIX B'): hp.paragraph_format.space_before=Pt(18)
        if t=="EXECUTIVE SUMMARY": ack=True
    elif l.startswith("@table"):
        parts=l[len("@table"):].strip().split("|"); title=parts[0]; hdr=parts[1:]
        rows=[hdr]
        while i<len(lines) and lines[i].strip() and not lines[i].startswith(("#","@","- ")) and "|" in lines[i]:
            rows.append(lines[i].split("|")); i+=1
        add_table(title,rows)
    elif l.startswith("@sig"):
        a=l[4:].strip().split("|")
        for k,x in enumerate(a):
            if x: para(x,12,bold=(k==0),align=WD_ALIGN_PARAGRAPH.RIGHT,spacing=1.0,after=2)
        pagebreak(); toc(); toc_done=True
    elif l.startswith("- "):
        p=para("",indent=0.9); p.paragraph_format.first_line_indent=Cm(-0.5)
        setfont(p.add_run("•  "+l[2:]),12)
    elif re.match(r"\d+\. ",l):
        p=para("",indent=0.9); p.paragraph_format.first_line_indent=Cm(-0.6)
        setfont(p.add_run(l),12)
    else:
        p=para(l); p.paragraph_format.first_line_indent=Cm(1.0) if False else None
d.save("Project_Report.docx")
